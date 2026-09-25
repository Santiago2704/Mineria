"""Actualiza el CONTROL REVOLVENTES con un nuevo "Informe de Vencimientos de Credito" del banco.

Uso:
    python3 actualizar_control.py CONTROL_ANTERIOR.xlsx INFORME_NUEVO.xls -o "CONTROL REVOLVENTES - <periodo>.xlsx"

- Reconstruye: RESUMEN, VENCEN - ACCION, Vigentes (cap=0) y una hoja por ejecutivo.
- Conserva intactas: CONTACTOS EMPRESAS, NOTAS - *, EXTENSIONES y cualquier otra hoja.
- Conserva lo capturado a mano (Intereses pagados, Forma de renovacion, Constancia
  elaborada, Estatus, Notas) cuando el mismo credito con la misma fecha vuelve a aparecer.
- Agrega a CONTACTOS EMPRESAS las empresas nuevas y llena la columna "Contacto" desde ahi.
- En las hojas de ejecutivo, arriba solo van los revolventes que VENCEN (capital a pagar != 0);
  los revolventes con capital 0 bajan a OTROS CREDITOS junto con simples, factoraje, etc.
"""
import argparse
import re
import unicodedata
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

EJECUTIVOS = [  # (nombre en el banco, nombre corto, pestaña)
    ("OSCAR ALEJANDRO MAYORGA MARTINEZ", "Oscar Mayorga", "OSCAR MARTINEZ"),
    ("HECTOR ELOY SALAZAR RIVERA", "Hector Salazar", "HECTOR RIVERA"),
    ("ELIAS FRANCISCO SIRENO QUINTERO", "Elias Sireno", "ELIAS QUINTERO"),
    ("FRANCISCO MARTINEZ BECERRA", "Francisco Martinez", "FRANCISCO BECERRA"),
]
GARANTIAS = [  # orden del resumen
    ("Liquida", "ALTA"), ("Fiduciaria", "ALTA"), ("Prendaria", "ALTA"),
    ("Hipotecaria", "MEDIA"), ("Deuda Solidaria / Fianza", "BAJA"),
]
DIAS = ["Lunes", "Martes", "Miercoles", "Jueves", "Viernes", "Sabado", "Domingo"]
MANUALES = ["Intereses pagados", "Forma de renovacion", "Constancia elaborada", "Estatus", "Notas"]
HEADERS = ["Vence", "Fecha Venc.", "Ejecutivo", "Cliente", "Contacto", "Producto", "Garantia", "Carga",
           "Divisa", "No. Cuenta", "No. Credito", "Capital Vigente", "Capital a Pagar (cuota)", "Interes",
           "IVA Interes", "Total a Pagar", "Plazo", "Tasa", "Fecha Inicio"] + MANUALES
WIDTHS = [9, 12, 32, 45, 28, 42, 24, 8, 7, 15, 16, 16, 18, 14, 13, 16, 7, 8, 12, 16, 20, 18, 14, 40]
REBUILT = ["RESUMEN", "VENCEN - ACCION", "Vigentes (cap=0)"] + [e[2] for e in EJECUTIVOS]

AZUL, AZUL2 = "1F4E79", "203864"
F_HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_BODY = Font(name="Calibri", size=10)
F_BOLD = Font(name="Calibri", size=10, bold=True)
F_TITLE = Font(name="Calibri", size=13, bold=True, color=AZUL)
F_SMALL = Font(name="Calibri", size=9)
FILL = {k: PatternFill("solid", fgColor=v) for k, v in {
    "hdr": AZUL, "sec": AZUL2, "rojo": "FFC7CE", "amarillo": "FFEB9C", "verde": "C6EFCE",
    "gris": "D9D9D9", "fecha": "FFFF00"}.items()}
CARGA_FILL = {"ALTA": FILL["rojo"], "MEDIA": FILL["amarillo"]}
THIN = Side("thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).upper().strip()


def nombre_bonito(cliente):
    t = str(cliente).title()
    for mal, bien in [(r"\bS\.?A\.? De C\.?V\.?", "SA de CV"), (r"\bSapi De Cv\b", "SAPI de CV"),
                      (r"\bS De Rl De Cv\b", "S de RL de CV"), (r"\bSofom Enr\b", "SOFOM ENR"),
                      (r"\bA\.C\.", "A.C."), (r"\bSc\b", "SC"), (r"\bSa\b", "SA"),
                      (r"\b(De|Del|La|Las|Los|Y|En|E)\b", lambda m: m.group(1).lower())]:
        t = re.sub(mal, bien, t, flags=re.I)
    return t[0].upper() + t[1:]


def es_revolvente(producto):
    p = norm(producto)
    return "C.C." in p or "CUENTA CORRIENTE" in p or "CTA CORRIENTE" in p


def clasificar(producto):
    p = norm(producto)
    if "LIQU" in p:
        return "Liquida", "ALTA"
    if "FIDU" in p:
        return "Fiduciaria", "ALTA"
    if "HIPOTEC" in p:
        return "Hipotecaria", "MEDIA"
    if "PRENDA" in p:
        return "Prendaria", "ALTA"
    if "DEUDA SOL" in p or "FIANZ" in p:
        return "Deuda Solidaria / Fianza", "BAJA"
    return "Sin clasificar", "REVISAR"


def to_num(v):
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return 0.0


def to_date(v):
    if isinstance(v, datetime):
        return v
    try:
        return datetime.strptime(str(v).strip()[:10], "%d/%m/%Y")
    except ValueError:
        return None


def leer_informe(path):
    raw = pd.read_excel(path, header=None, dtype=str)
    hdr = next(i for i, r in raw.iterrows() if str(r[0]).strip() == "Ejecutivo")
    encabezado = " ".join(str(x) for x in raw.iloc[:hdr, 0] if str(x) != "nan")
    df = pd.read_excel(path, header=hdr, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    nombres = {e[0]: e for e in EJECUTIVOS}
    filas = []
    for _, r in df.iterrows():
        ej = norm(r["Ejecutivo"])
        if ej not in nombres:
            continue
        gar, carga = clasificar(r["Producto"])
        interes, iva = to_num(r["Importe de Interés Proyectado"]), to_num(r["Importe de IVA Interés Proyectado"])
        cap = to_num(r["Importe de Capital a Pagar de la CUOTA"])
        filas.append({
            "ejecutivo": ej, "corto": nombres[ej][1], "cliente": str(r["Nombre del Cliente"]).strip(),
            "producto": str(r["Producto"]).strip(), "garantia": gar, "carga": carga,
            "divisa": str(r["Divisa"]).strip(), "cuenta": str(r["Numero de Cuenta"]).strip(),
            "credito": str(r["Número de Crédito"]).strip(), "cap_vig": to_num(r["Capital Vigente"]),
            "cap": cap, "interes": interes, "iva": iva, "total": to_num(r["Total a pagar"]),
            "plazo": int(to_num(r["Plazo Días"])), "tasa": to_num(r["Tasa de Interés"]),
            "inicio": to_date(r["Fecha de Inicio"]), "fecha": to_date(r["Fecha de Vencimiento"]),
            "revolvente": es_revolvente(r["Producto"]), "vence": cap != 0,
        })
    filas.sort(key=lambda f: (f["fecha"], f["ejecutivo"], f["cliente"], f["credito"]))
    return filas, encabezado


def leer_manuales(wb):
    """(credito, fecha) -> {columna manual: valor} de las hojas de detalle del control anterior."""
    datos = {}
    for ws in wb.worksheets:
        if ws.title not in REBUILT or ws.title == "RESUMEN":
            continue
        cols = {str(c.value).strip(): c.column for c in ws[1] if c.value}
        if "No. Credito" not in cols:
            continue
        for row in ws.iter_rows(min_row=2):
            cred = row[cols["No. Credito"] - 1].value
            fecha = to_date(row[cols["Fecha Venc."] - 1].value) if "Fecha Venc." in cols else None
            if not cred or not fecha:
                continue
            vals = {m: row[cols[m] - 1].value for m in MANUALES if m in cols}
            if any(v not in (None, "", "Pendiente", "No aplica") for v in vals.values()):
                datos[(str(cred).strip(), fecha.date())] = vals
    return datos


def actualizar_contactos(wb, filas):
    """Agrega empresas nuevas al final de CONTACTOS EMPRESAS y regresa cliente -> nombre del contacto."""
    ws = wb["CONTACTOS EMPRESAS"]
    last, existentes, contactos = 4, set(), {}
    for r in range(5, ws.max_row + 1):
        emp = ws.cell(r, 2).value
        if isinstance(emp, str) and emp.strip():
            last = r
            existentes.add(norm(emp))
            nombre = ws.cell(r, 4).value
            if nombre:
                contactos[norm(emp)] = " / ".join(x.strip() for x in str(nombre).splitlines() if x.strip())
    nuevas = {}
    for f in filas:
        if norm(f["cliente"]) not in existentes:
            nuevas.setdefault(norm(f["cliente"]), (f["cliente"], f["ejecutivo"]))
    borde = Border(*[Side("thin", color="8FAADC")] * 4)
    amarillo, alterno = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="EAF2FA")
    num = (ws.cell(last, 1).value or last - 4) if last > 4 else 0
    for i, (emp, ej) in enumerate(sorted(nuevas.values())):
        r = last + 1 + i
        banda = alterno if (r - 5) % 2 == 1 else PatternFill()
        for col in range(1, 12):
            c = ws.cell(r, col)
            c.font, c.border = F_BODY, borde
            c.alignment = Alignment(horizontal="center" if col in (1, 10) else "left", vertical="center", wrap_text=True)
            c.fill = banda if col in (1, 2, 3) else amarillo
        ws.cell(r, 1, int(num) + 1 + i)
        ws.cell(r, 2, emp)
        ws.cell(r, 3, ej)
        ws.cell(r, 10, "Pendiente")
    return contactos, [v[0] for v in sorted(nuevas.values())]


def hoja_detalle(wb, idx, titulo, secciones, contactos, manuales, freeze):
    ws = wb.create_sheet(titulo, idx)
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(1, i, h if i > 1 or titulo != "VENCEN - ACCION" else " ")
        c.font, c.fill, c.alignment = F_HDR, FILL["hdr"], CENTER
        ws.column_dimensions[c.column_letter].width = WIDTHS[i - 1]
    ws.row_dimensions[1].height = 32
    r = 2
    for nombre_sec, filas in secciones:
        if nombre_sec:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEADERS))
            c = ws.cell(r, 1, f"{nombre_sec}  ({len(filas)})")
            c.font, c.fill = Font(name="Calibri", size=11, bold=True, color="FFFFFF"), FILL["sec"]
            c.alignment = Alignment(horizontal="left", vertical="center")
            r += 1
            if not filas:
                ws.cell(r, 1, "Sin creditos en este periodo").font = Font(name="Calibri", size=10, italic=True, color="808080")
                r += 1
        for f in filas:
            man = manuales.get((f["credito"], f["fecha"].date()), {})
            vals = ["VENCE" if f["vence"] else "vigente", f["fecha"], f["ejecutivo"], f["cliente"],
                    contactos.get(norm(f["cliente"]), ""), f["producto"], f["garantia"], f["carga"],
                    f["divisa"], f["cuenta"], f["credito"], f["cap_vig"], f["cap"], f["interes"], f["iva"],
                    f["total"], f["plazo"], f["tasa"], f["inicio"],
                    man.get("Intereses pagados"), man.get("Forma de renovacion"), man.get("Constancia elaborada"),
                    man.get("Estatus") or ("Pendiente" if f["vence"] else "No aplica"), man.get("Notas")]
            for i, v in enumerate(vals, 1):
                c = ws.cell(r, i, v)
                c.font, c.border = F_BODY, BORDER
                if i in (2, 19):
                    c.number_format, c.alignment = "dd/mm/yyyy", Alignment(horizontal="center")
                elif 12 <= i <= 16:
                    c.number_format = "#,##0.00"
                elif i == 18:
                    c.number_format = "0.00"
            a = ws.cell(r, 1)
            if f["vence"]:
                a.font, a.fill = F_BOLD, (FILL["rojo"] if f["carga"] == "ALTA" else FILL["amarillo"])
            else:
                a.fill = FILL["verde"]
            if f["garantia"] == "Sin clasificar":
                ws.cell(r, 7).fill = FILL["amarillo"]
            r += 1
    fin = max(r - 1, 2)
    for formula, rango in [('"Si,No,N/A"', f"T2:T{fin} V2:V{fin}"),
                           ('"Banca electronica,Constancia,No renueva,Por definir"', f"U2:U{fin}"),
                           ('"Pendiente,En proceso,Listo,No aplica"', f"W2:W{fin}")]:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, sqref=rango)
        ws.add_data_validation(dv)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:X{fin}"
    return ws


def hoja_resumen(wb, filas, periodo, aviso):
    ws = wb.create_sheet("RESUMEN", 0)
    for col, w in zip("ABCDEFGHIJK", [14, 12, 20, 44, 15, 16, 24, 9, 16, 17, 8]):
        ws.column_dimensions[col].width = w
    ws["A1"] = "CONTROL DE VENCIMIENTOS - CREDITOS REVOLVENTES (CUENTA CORRIENTE)"
    ws["A1"].font = F_TITLE
    ws["A2"], ws["A3"] = periodo, f"Generado: {datetime.now():%d/%m/%Y %H:%M}"
    ws["A2"].font = ws["A3"].font = F_SMALL
    r = 4
    if aviso:
        ws.cell(r, 1, aviso).font = Font(name="Calibri", size=9, bold=True, color="C00000")
    r = 5

    def hdr(row, textos, cols):
        for col, t in zip(cols, textos):
            c = ws.cell(row, col, t)
            c.font, c.fill, c.alignment = F_HDR, FILL["hdr"], CENTER

    def celda(row, col, v, bold=False, fill=None, center=True, fmt=None):
        c = ws.cell(row, col, v)
        c.font = F_BOLD if bold else F_BODY
        if fill:
            c.fill = fill
        if center:
            c.alignment = Alignment(horizontal="center", vertical="center")
        if fmt:
            c.number_format = fmt
        return c

    # 1. Por ejecutivo (sin montos)
    ws.cell(r, 1, "Por ejecutivo").font = F_TITLE
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    hdr(r, ["Ejecutivo", "Revolventes", "Vencen", "Vigentes", "Otros creditos"], [1, 3, 4, 5, 6])
    tot = [0, 0, 0, 0]
    for ej, _, _ in EJECUTIVOS:
        r += 1
        mias = [f for f in filas if f["ejecutivo"] == ej]
        rev = [f for f in mias if f["revolvente"]]
        vals = [len(rev), sum(f["vence"] for f in rev), sum(not f["vence"] for f in rev), len(mias) - len(rev)]
        tot = [a + b for a, b in zip(tot, vals)]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        celda(r, 1, ej.title(), center=False)
        for i, v in enumerate(vals):
            celda(r, 3 + i, v, bold=(i == 1), fill=FILL["amarillo"] if i == 1 and v else None)
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    celda(r, 1, "TOTAL", bold=True, fill=FILL["gris"], center=False)
    for i, v in enumerate(tot):
        celda(r, 3 + i, v, bold=True, fill=FILL["gris"])

    # 2. Garantia x ejecutivo (solo los que vencen)
    vencen = [f for f in filas if f["revolvente"] and f["vence"]]
    r += 3
    ws.cell(r, 1, "Vencimientos por tipo de garantia y ejecutivo (mas garantia = mas trabajo)").font = F_TITLE
    r += 1
    hdr(r, ["Garantia", "Carga"] + [e[1] for e in EJECUTIVOS] + ["Total"], range(1, 8))
    presentes = [g for g in GARANTIAS if any(f["garantia"] == g[0] for f in vencen)]
    presentes += sorted({(f["garantia"], f["carga"]) for f in vencen} - set(GARANTIAS))
    if not presentes:
        r += 1
        ws.cell(r, 1, "Sin vencimientos de revolventes en este periodo").font = F_BODY
    for gar, carga in presentes:
        r += 1
        celda(r, 1, gar, bold=carga == "ALTA", center=False)
        celda(r, 2, carga, fill=CARGA_FILL.get(carga))
        for i, (ej, _, _) in enumerate(EJECUTIVOS):
            n = sum(f["ejecutivo"] == ej and f["garantia"] == gar for f in vencen)
            celda(r, 3 + i, n or "-", bold=bool(n))
        celda(r, 7, sum(f["garantia"] == gar for f in vencen), bold=True, fill=FILL["gris"])
    r += 1
    celda(r, 1, "TOTAL", bold=True, fill=FILL["gris"], center=False)
    celda(r, 2, None, fill=FILL["gris"])
    for i, (ej, _, _) in enumerate(EJECUTIVOS):
        celda(r, 3 + i, sum(f["ejecutivo"] == ej for f in vencen), bold=True, fill=FILL["gris"])
    celda(r, 7, len(vencen), bold=True, fill=FILL["gris"])

    # 3. Calendario, un renglon por credito: primero interes, luego capital
    r += 3
    ws.cell(r, 1, "Calendario de vencimientos - un renglon por credito").font = F_TITLE
    ws.cell(r, 4, "Primero se fondea el interes, luego se renueva. Por eso el interes va antes que el capital.").font = \
        Font(name="Calibri", size=9, color="808080")
    r += 1
    hdr(r, ["Fecha", "Dia", "Ejecutivo", "Cliente", "No. Cuenta", "No. Credito", "Garantia", "Carga",
            "1o INTERES a fondear", "2o Capital a pagar", "Divisa"], range(1, 12))
    fechas = sorted({f["fecha"] for f in vencen})
    for fecha in fechas:
        del_dia = [f for f in vencen if f["fecha"] == fecha]
        inicio = r + 1
        for f in del_dia:
            r += 1
            vals = [fecha.strftime("%d/%m/%Y"), DIAS[fecha.weekday()], f["corto"], nombre_bonito(f["cliente"]),
                    f["cuenta"], f["credito"], f["garantia"], f["carga"], f["interes"] + f["iva"], f["cap"], f["divisa"]]
            for i, v in enumerate(vals, 1):
                c = celda(r, i, v, center=i in (1, 2, 5, 6, 8, 11), fmt="#,##0.00" if i in (9, 10) else None)
                c.border = BORDER
            ws.cell(r, 1).font, ws.cell(r, 1).fill = F_BOLD, FILL["fecha"]
            if f["carga"] in CARGA_FILL:
                ws.cell(r, 8).fill = CARGA_FILL[f["carga"]]
            ws.cell(r, 9).font = Font(name="Calibri", size=10, bold=True, color="1F6E43")
        if r > inicio:
            for col in (1, 2):
                ws.merge_cells(start_row=inicio, start_column=col, end_row=r, end_column=col)
    if not fechas:
        r += 1
        ws.cell(r, 1, "Sin vencimientos de revolventes en este periodo").font = F_BODY

    # 4. Recordatorio del proceso
    r += 3
    ws.cell(r, 1, "Recordatorio del proceso").font = F_TITLE
    for linea in ["1. El cliente paga PRIMERO los intereses el dia del vencimiento.",
                  "2. Solo despues de pagar intereses puede renovar.",
                  "3. Renovacion por banca electronica -> tramite minimo de tu parte.",
                  "4. Renovacion por CONSTANCIA -> tu elaboras el documento.",
                  "5. Garantia liquida, fiduciaria o prendaria -> carga de trabajo ALTA.",
                  "6. Capital a pagar = 0 -> aun no se cumplen los 90 dias, no hay nada que hacer.",
                  "7. El proceso empieza 2 dias antes del vencimiento."]:
        r += 1
        ws.cell(r, 1, linea).font = F_BODY
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("control")
    ap.add_argument("informe")
    ap.add_argument("-o", "--output", required=True)
    args = ap.parse_args()

    filas, encabezado = leer_informe(args.informe)
    wb = load_workbook(args.control)
    manuales = leer_manuales(wb)
    contactos, nuevas = actualizar_contactos(wb, filas)

    orden = [ws.title for ws in wb.worksheets]
    posiciones = {t: orden.index(t) for t in REBUILT if t in orden}
    for t in REBUILT:
        if t in wb.sheetnames:
            del wb[t]

    ini, fin = min(f["fecha"] for f in filas), max(f["fecha"] for f in filas)
    periodo = f"Vencimientos del {ini:%d/%m/%Y} al {fin:%d/%m/%Y}"
    fechas_hdr = [to_date(x) for x in re.findall(r"\d{2}/\d{2}/\d{4}", encabezado)]
    aviso = None
    if len(fechas_hdr) == 2 and not (fechas_hdr[0] <= ini and fin <= fechas_hdr[1]):
        aviso = (f'OJO: el encabezado del banco dice "{encabezado.strip()}", '
                 "pero los datos no caen en ese rango. Se uso el rango real de los datos.")

    hoja_resumen(wb, filas, periodo, aviso)
    rev = [f for f in filas if f["revolvente"]]
    hoja_detalle(wb, posiciones.get("VENCEN - ACCION", 2), "VENCEN - ACCION",
                 [(None, [f for f in rev if f["vence"]])], contactos, manuales, "F2")
    hoja_detalle(wb, posiciones.get("Vigentes (cap=0)", 3), "Vigentes (cap=0)",
                 [(None, [f for f in rev if not f["vence"]])], contactos, manuales, "F2")
    for i, (ej, _, tab) in enumerate(EJECUTIVOS):
        mias = [f for f in filas if f["ejecutivo"] == ej]
        arriba = [f for f in mias if f["revolvente"] and f["vence"]]
        abajo = ([f for f in mias if f["revolvente"] and not f["vence"]]
                 + [f for f in mias if not f["revolvente"]])
        hoja_detalle(wb, posiciones.get(tab, 4 + i), tab, [
            ("REVOLVENTES QUE VENCEN - Cuenta Corriente / C.C. con capital a pagar  (los importantes)", arriba),
            ("OTROS CREDITOS - revolventes sin capital a pagar (cap=0), simples, arrendamiento, factoraje, etc.", abajo),
        ], contactos, manuales, "L2")
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "RESUMEN"
    wb.save(args.output)

    # Resumen en consola para verificar
    print(periodo)
    print(f"Creditos de los 4 ejecutivos: {len(filas)} | revolventes: {len(rev)} | vencen: {sum(f['vence'] for f in rev)}")
    for ej, corto, _ in EJECUTIVOS:
        r_ = [f for f in rev if f["ejecutivo"] == ej]
        print(f"  {corto}: {len(r_)} revolventes, {sum(f['vence'] for f in r_)} vencen")
    sin = sorted({f["producto"] for f in rev if f["garantia"] == "Sin clasificar"})
    if sin:
        print("REVISAR - revolventes sin clasificar:", sin)
    print(f"Capturas manuales conservadas: {sum((f['credito'], f['fecha'].date()) in manuales for f in filas)}")
    print(f"Empresas nuevas en CONTACTOS: {len(nuevas)}", *nuevas, sep="\n  ")


if __name__ == "__main__":
    main()
