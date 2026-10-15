import sys, re, openpyxl
from collections import OrderedDict
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
ws = wb['PromoterComissions']
for n in ('Base de Datos', 'Resumen por Promotor'):
    if n in wb.sheetnames:
        del wb[n]

def num(v):
    if v in (None, ''):
        return None
    if isinstance(v, (int, float)):
        return v
    return float(str(v).replace(',', ''))

periodo = None
rows = []
prom = emp = port = div = None
for r in range(1, ws.max_row + 1):
    a = ws.cell(r, 1).value
    if isinstance(a, str) and 'Reporte de Operaciones' in a:
        m = re.search(r'del (\S+) al (\S+)', a)
        periodo = f'{m.group(1)} al {m.group(2)}'
    b = ws.cell(r, 2).value
    if b == 'Promotor:':
        prom, emp = ws.cell(r, 6).value, ws.cell(r, 39).value
    elif b == 'Portafolio:':
        port = ws.cell(r, 8).value
    elif b == 'Divisa:':
        div = ws.cell(r, 4).value
    elif isinstance(b, int):
        g = lambda col: ws[f'{col}{r}'].value
        rows.append([prom, emp, port, div, b, g('D'), g('J'), g('M'), g('U'), g('W'),
                     num(g('Y')), num(g('AC')), num(g('AI')), num(g('AK')),
                     num(g('AP')), num(g('AT')), num(g('AY'))])

hdr_font = Font(bold=True, color='FFFFFF', size=11)
hdr_fill = PatternFill('solid', fgColor='1F4E78')
thin = Side(style='thin')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
f10 = Font(size=10)
white = PatternFill('solid', fgColor='FFFFFF')

# Base de Datos
bd = wb.create_sheet('Base de Datos')
cols = ['Promotor', 'No. Empleado', 'Portafolio', 'Divisa', 'No. Operación', 'Fecha Operación',
        'Fecha Valor', 'Cliente', 'Instrumento', '# Liq.', 'Compra', 'Venta', 'Tipo de Cambio',
        'T.C. Mesa', 'Utilidad Bruta', 'Utilidad Neta', 'Costo Producto']
widths = [35, 15, 28, 18, 15, 16, 12, 55, 14, 8, 16, 16, 15, 15, 16, 14, 14]
fmts = {11: '#,##0.00', 12: '#,##0.00', 13: '#,##0.0000', 14: '#,##0.0000',
        15: '#,##0.00', 16: '#,##0.00', 17: '#,##0.000000'}
for i, (h, w) in enumerate(zip(cols, widths), 1):
    c = bd.cell(1, i, h)
    c.font, c.fill, c.border, c.alignment = hdr_font, hdr_fill, box, center
    bd.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
bd.row_dimensions[1].height = 20
for ri, row in enumerate(rows, 2):
    for ci, v in enumerate(row, 1):
        c = bd.cell(ri, ci, v)
        c.font, c.fill, c.border = f10, white, box
        if ci in fmts:
            c.number_format = fmts[ci]

# Resumen
rs = wb.create_sheet('Resumen por Promotor')
rs['A1'] = 'Resumen de Utilidad Bruta por Promotor / Portafolio / Divisa'
rs['A1'].font = Font(bold=True, size=12)
rs['A2'] = f'Periodo: {periodo}'
for i, h in enumerate(['Promotor', 'Portafolio', 'Divisa', '# Operaciones', 'Utilidad Bruta', '% del Total'], 1):
    c = rs.cell(4, i, h)
    c.font, c.fill, c.border, c.alignment = hdr_font, hdr_fill, box, center
for col, w in zip('ABCDEFGHI', [42, 28, 18, 14, 18, 14, 4, 28, 18]):
    rs.column_dimensions[col].width = w

total = sum(r[14] or 0 for r in rows)
grp = OrderedDict()
for r in rows:
    k = grp.setdefault(r[0], OrderedDict()).setdefault((r[2], r[3]), [0, 0])
    k[0] += 1
    k[1] += r[14] or 0

sub_font = Font(bold=True, color='1F4E78', size=10)
sub_fill = PatternFill('solid', fgColor='D9E1F2')
tot_font = Font(bold=True, color='FFFFFF')
r = 5
for p, d in grp.items():
    for (po, dv), (n, ub) in d.items():
        for ci, v in enumerate([p, po, dv, n, ub, ub / total if total else 0], 1):
            c = rs.cell(r, ci, v)
            c.font = f10
        rs.cell(r, 4).number_format = '#,##0'
        rs.cell(r, 5).number_format = '#,##0.00'
        rs.cell(r, 6).number_format = '0.00%'
        r += 1
    n = sum(x[0] for x in d.values()); ub = sum(x[1] for x in d.values())
    for ci, v in enumerate([f'Subtotal {p}', None, None, n, ub, ub / total if total else 0], 1):
        c = rs.cell(r, ci, v)
        c.font, c.fill = sub_font, sub_fill
    rs.cell(r, 4).number_format = '#,##0'
    rs.cell(r, 5).number_format = '#,##0.00'
    rs.cell(r, 6).number_format = '0.00%'
    r += 2
for ci, v in enumerate(['TOTAL GENERAL', None, None, len(rows), total, 1], 1):
    c = rs.cell(r, ci, v)
    c.font, c.fill = tot_font, hdr_fill
rs.cell(r, 4).number_format = '#,##0'
rs.cell(r, 5).number_format = '#,##0.00'
rs.cell(r, 6).number_format = '0.00%'

# Totales por portafolio (lateral)
pt = {}
for x in rows:
    pt[x[2]] = pt.get(x[2], 0) + (x[14] or 0)
rr = 5
for po in sorted(pt):
    rs.cell(rr, 8, po).font = Font(bold=True, size=10)
    c = rs.cell(rr, 9, pt[po]); c.font = f10; c.number_format = '#,##0.00'
    rr += 1
rs.cell(rr, 8, 'Total').font = Font(bold=True, size=10)
c = rs.cell(rr, 9, total); c.font = Font(bold=True, size=10); c.number_format = '#,##0.00'

wb.save(out)
print(f'{len(rows)} operaciones, {len(grp)} promotores, UB total {total:,.2f}, periodo {periodo}')
